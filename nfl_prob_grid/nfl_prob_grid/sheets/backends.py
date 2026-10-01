"""Spreadsheet backends with one small interface.

* ``GoogleSheetsBackend`` - the live Google Sheet via the Sheets API (gspread + service account).
* ``XlsxBackend``         - a local .xlsx with the same tab layout (offline use and tests).

Both expose: tabs(), read(tab) -> list of rows, write(TabSpec) (create or fully replace a tab), flush().
Reads return raw values: numbers as numbers, dates/times as datetimes (xlsx) or spreadsheet serial numbers
(Google, UNFORMATTED/SERIAL_NUMBER rendering) - the bridge normalises both. Writes are values only (RAW):
nothing the bridge writes can be re-interpreted as a formula.
"""
from __future__ import annotations

import math
import os
import re
import tempfile
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path


@dataclass
class TabSpec:
    """A fully specified output tab. Row/column indices are 0-based."""
    name: str
    values: list                                   # 2-D list of str/float/int/None
    freeze_rows: int = 1
    freeze_cols: int = 0
    bold_rows: set = field(default_factory=set)
    number_formats: dict = field(default_factory=dict)   # (r0, c0, r1, c1) inclusive -> "0.0%" etc.
    fills: dict = field(default_factory=dict)            # (r, c) -> "#RRGGBB"
    col_widths: dict = field(default_factory=dict)       # c -> pixels
    tab_color: str = "#1F4E79"


def clean(v):
    """JSON/xlsx-safe cell value."""
    if v is None:
        return ""
    if hasattr(v, "item") and not isinstance(v, (str, bytes)):
        v = v.item()                                   # numpy scalar -> python
    if isinstance(v, float) and (math.isnan(v) or math.isinf(v)):
        return ""
    if isinstance(v, (datetime, date)):
        return v.isoformat(sep=" ") if isinstance(v, datetime) else v.isoformat()
    return v


def _hex_rgb(h: str) -> dict:
    h = h.lstrip("#")
    return {"red": int(h[0:2], 16) / 255, "green": int(h[2:4], 16) / 255, "blue": int(h[4:6], 16) / 255}


# =====================================================================================================
# Google Sheets
# =====================================================================================================
class GoogleSheetsBackend:
    def __init__(self, spreadsheet: str, credentials: str, client=None):
        if client is None:
            try:
                import gspread
            except ImportError as exc:  # pragma: no cover
                raise RuntimeError("Google backend needs: pip install gspread google-auth") from exc
            if not Path(credentials).exists():
                raise RuntimeError(f"Service-account key not found: {credentials} (see SHEETS_SETUP.md)")
            client = gspread.service_account(filename=credentials)
        if re.match(r"^https?://", spreadsheet):
            self.sh = client.open_by_url(spreadsheet)
        else:
            self.sh = client.open_by_key(spreadsheet)

    def tabs(self) -> list[str]:
        return [ws.title for ws in self.sh.worksheets()]

    def read(self, tab: str) -> list[list]:
        ws = self.sh.worksheet(tab)
        return ws.get_all_values(value_render_option="UNFORMATTED_VALUE",
                                 date_time_render_option="SERIAL_NUMBER")

    def write(self, spec: TabSpec) -> None:
        rows = [[clean(v) for v in r] for r in spec.values]
        n_rows, n_cols = max(len(rows), 1), max((len(r) for r in rows), default=1)
        rows = [r + [""] * (n_cols - len(r)) for r in rows]
        try:
            ws = self.sh.worksheet(spec.name)
            ws.clear()
            self.sh.batch_update({"requests": [{"unmergeCells": {"range": {"sheetId": ws.id}}},
                                               {"repeatCell": {"range": {"sheetId": ws.id},
                                                               "cell": {"userEnteredFormat": {}},
                                                               "fields": "userEnteredFormat"}}]})
            if ws.row_count < n_rows or ws.col_count < n_cols:
                ws.resize(rows=max(ws.row_count, n_rows), cols=max(ws.col_count, n_cols))
        except Exception as exc:  # gspread.WorksheetNotFound
            if exc.__class__.__name__ != "WorksheetNotFound":
                raise
            ws = self.sh.add_worksheet(title=spec.name, rows=max(n_rows, 20), cols=max(n_cols, 5))
        ws.update(rows, "A1", value_input_option="RAW")
        self.sh.batch_update({"requests": self.format_requests(ws.id, spec, n_rows, n_cols)})

    @staticmethod
    def format_requests(sheet_id: int, spec: TabSpec, n_rows: int, n_cols: int) -> list[dict]:
        rq = [{"updateSheetProperties": {
            "properties": {"sheetId": sheet_id, "tabColorStyle": {"rgbColor": _hex_rgb(spec.tab_color)},
                           "gridProperties": {"frozenRowCount": spec.freeze_rows, "frozenColumnCount": spec.freeze_cols}},
            "fields": "tabColorStyle,gridProperties.frozenRowCount,gridProperties.frozenColumnCount"}},
            {"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": 0, "endRowIndex": n_rows,
                                      "startColumnIndex": 0, "endColumnIndex": n_cols},
                            "cell": {"userEnteredFormat": {"textFormat": {"fontFamily": "Arial", "fontSize": 10}}},
                            "fields": "userEnteredFormat.textFormat"}}]
        for r in sorted(spec.bold_rows):
            rq.append({"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": r, "endRowIndex": r + 1,
                                                "startColumnIndex": 0, "endColumnIndex": n_cols},
                                      "cell": {"userEnteredFormat": {"textFormat": {"bold": True, "fontFamily": "Arial"}}},
                                      "fields": "userEnteredFormat.textFormat"}})
        for (r0, c0, r1, c1), fmt in spec.number_formats.items():
            rq.append({"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": r0, "endRowIndex": r1 + 1,
                                                "startColumnIndex": c0, "endColumnIndex": c1 + 1},
                                      "cell": {"userEnteredFormat": {"numberFormat": {"type": "NUMBER", "pattern": fmt}}},
                                      "fields": "userEnteredFormat.numberFormat"}})
        # fills: compress horizontal runs of the same colour into one request each
        by_row: dict[int, list] = {}
        for (r, c), col in spec.fills.items():
            by_row.setdefault(r, []).append((c, col))
        for r, cells in sorted(by_row.items()):
            cells.sort()
            start, prev_c, prev_col = cells[0][0], cells[0][0], cells[0][1]
            for c, col in cells[1:] + [(None, None)]:
                if c is not None and c == prev_c + 1 and col == prev_col:
                    prev_c = c
                    continue
                rq.append({"repeatCell": {"range": {"sheetId": sheet_id, "startRowIndex": r, "endRowIndex": r + 1,
                                                    "startColumnIndex": start, "endColumnIndex": prev_c + 1},
                                          "cell": {"userEnteredFormat": {"backgroundColor": _hex_rgb(prev_col)}},
                                          "fields": "userEnteredFormat.backgroundColor"}})
                if c is not None:
                    start, prev_c, prev_col = c, c, col
        for c, px in spec.col_widths.items():
            rq.append({"updateDimensionProperties": {"range": {"sheetId": sheet_id, "dimension": "COLUMNS",
                                                               "startIndex": c, "endIndex": c + 1},
                                                     "properties": {"pixelSize": px}, "fields": "pixelSize"}})
        return rq

    def flush(self) -> None:
        pass                                            # every write is already committed server-side


# =====================================================================================================
# Local .xlsx (same layout as the Google Sheet)
# =====================================================================================================
class XlsxBackend:
    """Local .xlsx mode. The input workbook is opened READ-ONLY and never saved: any library that re-saves a
    workbook rewrites every tab (floats can shift in the 16th digit, Excel-only features can be dropped).
    Output tabs therefore go to a companion workbook, by default "<name> - Grid Outputs.xlsx" next to it."""

    def __init__(self, path: str | Path, out_path: str | Path | None = None):
        import openpyxl
        self.path = Path(path)
        if not self.path.exists():
            raise RuntimeError(f"Workbook not found: {self.path}")
        self.out_path = Path(out_path) if out_path else self.path.with_name(f"{self.path.stem} - Grid Outputs.xlsx")
        if self.out_path.resolve() == self.path.resolve():
            raise RuntimeError("Output workbook must differ from the input workbook")
        self.values = openpyxl.load_workbook(self.path, data_only=True, read_only=True)
        if self.out_path.exists():
            self.out = openpyxl.load_workbook(self.out_path)
        else:
            self.out = openpyxl.Workbook()
            self.out.remove(self.out.active)
        self.dirty = False

    def tabs(self) -> list[str]:
        return list(self.values.sheetnames) + [t for t in self.out.sheetnames if t not in self.values.sheetnames]

    def read(self, tab: str) -> list[list]:
        ws = self.values[tab] if tab in self.values.sheetnames else self.out[tab]
        return [list(row) for row in ws.iter_rows(values_only=True)]

    def write(self, spec: TabSpec) -> None:
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        if spec.name in self.values.sheetnames:
            raise RuntimeError(f"Refusing to shadow input tab {spec.name!r}")
        wb = self.out
        if spec.name in wb.sheetnames:
            idx = wb.sheetnames.index(spec.name)
            del wb[spec.name]
            ws = wb.create_sheet(spec.name, idx)
        else:
            ws = wb.create_sheet(spec.name)
        base, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
        for r, row in enumerate(spec.values):
            for c, v in enumerate(row):
                v = clean(v)
                if v == "":
                    continue
                cell = ws.cell(r + 1, c + 1, v)
                if isinstance(v, str) and v.startswith("="):
                    cell.data_type = "s"                               # never a formula
                cell.font = bold if r in spec.bold_rows else base
        for (r0, c0, r1, c1), fmt in spec.number_formats.items():
            for r in range(r0, r1 + 1):
                for c in range(c0, c1 + 1):
                    ws.cell(r + 1, c + 1).number_format = fmt
        for (r, c), col in spec.fills.items():
            ws.cell(r + 1, c + 1).fill = PatternFill("solid", fgColor=col.lstrip("#"))
        for c, px in spec.col_widths.items():
            ws.column_dimensions[get_column_letter(c + 1)].width = max(px / 7, 6)
        if spec.freeze_rows or spec.freeze_cols:
            ws.freeze_panes = ws.cell(spec.freeze_rows + 1, spec.freeze_cols + 1)
        ws.sheet_properties.tabColor = spec.tab_color.lstrip("#")
        self.dirty = True

    def flush(self) -> None:
        if not self.dirty:
            return
        fd, tmp = tempfile.mkstemp(dir=self.out_path.parent, suffix=".xlsx")
        os.close(fd)
        try:
            self.out.save(tmp)
            os.replace(tmp, self.out_path)
        finally:
            if os.path.exists(tmp):
                os.unlink(tmp)
        self.dirty = False


def open_backend(spreadsheet: str, credentials: str):
    if not spreadsheet:
        raise RuntimeError("No spreadsheet configured: set [sheets] spreadsheet in config.toml or pass --spreadsheet")
    if spreadsheet.lower().endswith((".xlsx", ".xlsm")):
        return XlsxBackend(spreadsheet)
    return GoogleSheetsBackend(spreadsheet, credentials)
