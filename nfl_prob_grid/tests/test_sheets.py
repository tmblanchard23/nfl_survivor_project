"""Sheets bridge tests: the user's real workbook layout (xlsx backend), an in-memory fake of the Google backend,
and the exact Sheets API requests the Google backend would send."""
import shutil
import sys
from pathlib import Path

import openpyxl
import pandas as pd
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from nfl_prob_grid import load_config  # noqa: E402
from nfl_prob_grid.errors import ConfigError, InputError, StateError  # noqa: E402
from nfl_prob_grid.master import StateStore  # noqa: E402
from nfl_prob_grid.sheets import bootstrap, build, publish, sheet_status  # noqa: E402
from nfl_prob_grid.sheets.backends import GoogleSheetsBackend, TabSpec, XlsxBackend  # noqa: E402
from nfl_prob_grid.sheets.bridge import read_lines  # noqa: E402
from nfl_prob_grid.teams import TeamNormalizer  # noqa: E402
from nfl_prob_grid.weeks import assign_weeks, infer_current_week  # noqa: E402

FIX = ROOT / "tests" / "fixtures" / "sheet_example.xlsx"
NOW = "2026-09-23T14:00:00Z"
INPUT_TABS = [  # (kept for reference)
    "Game Lines", "Preserved Week 1 Full Season Li", "Preseason Win Totals", "Pools Picks",
              "Preserved Week 2 Line Export", "Super Bowl Futures", "Recommendation"]


@pytest.fixture()
def sheet(tmp_path):
    p = tmp_path / "sheet.xlsx"
    shutil.copy(FIX, p)
    return p


@pytest.fixture()
def cfg(tmp_path, sheet):
    return load_config(ROOT / "config.toml", paths__state_dir=str(tmp_path / "state"),
                       paths__output_dir=str(tmp_path / "out"), sheets__spreadsheet=str(sheet))


def boot(cfg):
    return bootstrap(cfg, catch_up=["Preserved Week 2 Line Export"], as_of=NOW)


def snapshot(path, tabs):
    wb = openpyxl.load_workbook(path)
    return {t: [[c.value for c in row] for row in wb[t].iter_rows()] for t in tabs}


def outputs(sheet):
    return sheet.with_name(f"{sheet.stem} - Grid Outputs.xlsx")


# ---------------------------------------------------------------- reading the user's layout
def test_reads_real_layout_header_row_found_and_times_normalised(cfg, sheet):
    be = XlsxBackend(sheet)
    lines = read_lines(be, "Game Lines", cfg)                  # metadata rows 1-3, header on row 5
    assert len(lines) == 32 and "week" not in lines            # the lines tab has no week column
    assert lines.updated.iloc[0] == "2026-09-17 16:24:00" and lines.game_time.iloc[0] == "2026-09-17 20:15:00"
    init = read_lines(be, "Preserved Week 1 Full Season Li", cfg)
    assert len(init) == 272 and set(init.week.astype(int)) == set(range(1, 19))


def test_google_serial_numbers_are_converted_like_datetimes():
    from nfl_prob_grid.sheets.bridge import _cell_time
    # 2026-09-17 16:24 as a Sheets serial number (days since 1899-12-30)
    serial = (pd.Timestamp("2026-09-17 16:24") - pd.Timestamp("1899-12-30")) / pd.Timedelta(days=1)
    assert _cell_time(serial, "America/New_York") == "2026-09-17 16:24:00"
    assert _cell_time("", "America/New_York") == "" and _cell_time(None, "America/New_York") == ""


# ---------------------------------------------------------------- week inference (the ID system does the work)
def _master(cfg):
    return StateStore(cfg.paths.state_dir).load()[0]


def test_weeks_are_inferred_by_teams_and_nearest_date_including_division_rematches(cfg):
    boot(cfg)
    m = _master(cfg)
    nz = TeamNormalizer()
    pair = m[(m.home_team == "BUF") | (m.away_team == "BUF")]
    div = pair.groupby(pair[["home_team", "away_team"]].apply(lambda r: frozenset(r), axis=1)).filter(lambda g: len(g) == 2)
    g1, g2 = div.sort_values("week").iloc[0], div.sort_values("week").iloc[1]
    rows = pd.DataFrame({"away": ["Buffalo Bills" if g2.away_team == "BUF" else "x"] * 0 + [g2.away_team, g1.away_team],
                         "home": [g2.home_team, g1.home_team],
                         "game_time": [g2.game_date + " 13:00:00", g1.game_date + " 13:00:00"]})
    out, aside = assign_weeks(rows, m, nz, "America/New_York", 6)
    assert list(out.week) == [int(g2.week), int(g1.week)] and not aside


def test_unscheduled_and_far_dates_are_set_aside_but_unknown_teams_go_to_the_engine(cfg):
    boot(cfg)
    m = _master(cfg)
    g = m.iloc[40]
    rows = pd.DataFrame({"away": [g.away_team, g.away_team, "Gotham Rogues"],
                         "home": [g.home_team, g.home_team, "Buffalo Bills"],
                         "game_time": ["2027-01-20 20:00:00", "", "2026-10-04 13:00:00"]})
    out, aside = assign_weeks(rows, m, TeamNormalizer(), "America/New_York", 6)
    assert [a["reason"] for a in aside] == ["date_far_from_schedule"]   # e.g. a playoff game
    assert len(out) == 2 and out.week.iloc[1] == ""                      # unknown team kept -> engine rejects loudly


def test_current_week_freezes_the_day_after_the_last_game(cfg):
    boot(cfg)
    m = _master(cfg)
    last_w2 = pd.to_datetime(m[m.week == 2].game_date).max()             # Monday night game date
    on_monday = pd.Timestamp(last_w2.strftime("%Y-%m-%d") + " 23:30", tz="America/New_York")
    tuesday = on_monday + pd.Timedelta(hours=12)
    assert infer_current_week(m, on_monday, "America/New_York", 18) == 2
    assert infer_current_week(m, tuesday, "America/New_York", 18) == 3
    assert infer_current_week(m, pd.Timestamp("2027-03-01", tz="UTC"), "America/New_York", 18) == 19


# ---------------------------------------------------------------- bootstrap / build / publish on the real workbook
def test_bootstrap_builds_full_history_and_never_touches_input_tabs(cfg, sheet):
    before = sheet.read_bytes()
    reps = boot(cfg)
    assert [r.version for r in reps] == [1, 2, 3]
    assert reps[1].counts["lines_refreshed"] == 32 and not reps[1].rejected_rows      # weeks 2+3 matched by ID
    assert reps[2].details["bridge"]["completed_week_rows"] == 16                     # last week's games skipped quietly
    assert not [i for i in reps[2].issues if i.code == "frozen_week_line_ignored"]
    assert sheet.read_bytes() == before                                               # user's workbook: byte-identical
    tabs = openpyxl.load_workbook(outputs(sheet)).sheetnames
    for t in ("Probability Grid (Latest)", "Grid 2026 W01 v001", "Grid 2026 W02 v002", "Grid 2026 W03 v003",
              "Master Game Table", "Team Ratings", "Grid Run Log"):
        assert t in tabs, t


def test_grid_tab_matches_master_exactly_and_marks_frozen_weeks(cfg, sheet):
    boot(cfg)
    m = _master(cfg)
    ws = openpyxl.load_workbook(outputs(sheet))["Probability Grid (Latest)"]
    assert [ws.cell(1, c).value for c in (1, 2, 19)] == ["Team", "Week 1", "Week 18"]
    names = {ws.cell(r, 1).value: r for r in range(2, 34)}
    from nfl_prob_grid.sheets.render import FULL_NAME
    for g in m.itertuples():
        assert ws.cell(names[FULL_NAME[g.home_team]], g.week + 1).value == pytest.approx(g.home_win_prob, abs=1e-6)
        assert ws.cell(names[FULL_NAME[g.away_team]], g.week + 1).value == pytest.approx(g.away_win_prob, abs=1e-6)
    grey = [ws.cell(names["Buffalo Bills"], c).fill.fgColor.rgb for c in (2, 3, 4)]
    assert grey[0].endswith("D9D9D9") and grey[1].endswith("D9D9D9") and not grey[2].endswith("D9D9D9")
    assert ws.cell(2, 2).number_format == "0.0%"


def test_bootstrap_refuses_to_overwrite_state_and_force_archives(cfg):
    boot(cfg)
    with pytest.raises(StateError, match="already exists"):
        boot(cfg)
    assert bootstrap(cfg, as_of=NOW, force=True)[0].version == 1


def test_weekly_build_with_new_lines_and_rerun_is_idempotent(cfg, sheet):
    boot(cfg)
    wb = openpyxl.load_workbook(sheet)                                  # simulate next week's pull into Game Lines
    ws = wb["Game Lines"]
    m = _master(cfg)
    from nfl_prob_grid.sheets.render import FULL_NAME
    w4 = m[m.week == 4]
    for r in range(6, ws.max_row + 1):
        for c in range(1, 12):
            ws.cell(r, c).value = None
    for i, g in enumerate(w4.itertuples()):
        row = [pd.Timestamp(g.game_date + " 13:00"), FULL_NAME[g.away_team], FULL_NAME[g.home_team], 2.4, 1.6,
               None, None, None, None, pd.Timestamp("2026-09-29 09:00")]
        for c, v in enumerate(row, start=1):
            ws.cell(6 + i, c).value = v
    wb.save(sheet)
    rep = build(cfg, as_of="2026-09-30T15:00:00Z")      # preseason lines (9/8) are now > 3 weeks old
    assert rep.counts["lines_refreshed"] == len(w4) and rep.current_week == 4 and not rep.rejected_rows
    assert rep.counts["stale_adjusted"] > 0                              # preseason lines now > 3 weeks old
    rep2 = build(cfg, as_of="2026-09-30T16:00:00Z")                     # same pull again
    assert rep2.counts.get("lines_refreshed", 0) == 0 and rep2.counts["ignored_same_timestamp"] == len(w4)
    wb = openpyxl.load_workbook(outputs(sheet))
    assert "Grid 2026 W04 v004" in wb.sheetnames and "Grid 2026 W04 v005" in wb.sheetnames
    log = wb["Grid Run Log"]
    assert [log.cell(r, 1).value for r in range(2, log.max_row + 1)] == ["v001", "v002", "v003", "v004", "v005"]


def test_publish_is_idempotent_and_repairs_a_failed_sheet_write(cfg, sheet):
    boot(cfg)
    out = outputs(sheet)
    first = snapshot(out, ["Probability Grid (Latest)", "Grid Run Log", "Master Game Table"])
    wb = openpyxl.load_workbook(out)
    del wb["Probability Grid (Latest)"]                                  # as if the write had failed
    wb.save(out)
    publish(cfg)
    publish(cfg)
    assert snapshot(out, ["Probability Grid (Latest)", "Grid Run Log", "Master Game Table"]) == first


def test_dry_run_writes_nothing(cfg, sheet):
    boot(cfg)
    before_state = StateStore(cfg.paths.state_dir).read_manifest()
    before = outputs(sheet).read_bytes()
    build(cfg, as_of="2026-09-24T12:00:00Z", dry_run=True)
    assert outputs(sheet).read_bytes() == before and StateStore(cfg.paths.state_dir).read_manifest() == before_state


def test_output_tab_names_can_never_be_input_tabs(cfg):
    cfg.sheets.grid_tab = "Game Lines"
    with pytest.raises(ConfigError, match="refusing to overwrite inputs"):
        build(cfg)


def test_bad_rows_in_lines_tab_are_reported_not_fatal(cfg, sheet):
    boot(cfg)
    wb = openpyxl.load_workbook(sheet)
    ws = wb["Game Lines"]
    ws.cell(7, 2).value = "Gotham Rogues"          # a typo'd team
    ws.cell(8, 4).value = None                     # one side's odds missing
    ws.cell(8, 5).value = None                     # both missing -> "no line yet"
    for r in range(6, ws.max_row + 1):             # make the pull newer than the stored one
        if ws.cell(r, 10).value:
            ws.cell(r, 10).value = pd.Timestamp("2026-09-23 09:00")
    wb.save(sheet)
    rep = build(cfg, as_of=NOW)
    b = rep.details["bridge"]
    assert [r["reason"] for r in b["set_aside_rows"]] == ["no_line_yet"]
    assert [r["code"] for r in rep.rejected_rows] == ["unknown_team"]                 # the typo, clearly named
    assert "Gotham Rogues" in rep.rejected_rows[0]["detail"]
    assert rep.counts["lines_refreshed"] > 0                                           # good rows still applied


def test_status_reports_inputs_and_state(cfg):
    s = sheet_status(cfg)
    assert s["engine_version"] is None and all(s["tabs_found"].values()) and s["lines_tab_rows"] == 32
    boot(cfg)
    s = sheet_status(cfg)
    assert s["engine_version"] == 3 and "Probability Grid (Latest)" in s["output_tabs_present"]


def test_missing_input_tab_is_a_clear_error(cfg):
    boot(cfg)
    cfg.sheets.lines_tab = "Lines That Do Not Exist"
    assert "not found" in sheet_status(cfg)["lines_tab_error"]
    with pytest.raises(InputError, match="not found"):
        build(cfg, as_of=NOW)


# ---------------------------------------------------------------- Google backend (fake gspread client)
class FakeWS:
    def __init__(self, title, sid, values=None):
        self.title, self.id, self.values, self.row_count, self.col_count = title, sid, values or [], 100, 26
        self.cleared = 0

    def get_all_values(self, **kw):
        assert kw == {"value_render_option": "UNFORMATTED_VALUE", "date_time_render_option": "SERIAL_NUMBER"}
        return self.values

    def clear(self):
        self.values, self.cleared = [], self.cleared + 1

    def update(self, rows, rng, value_input_option):
        assert rng == "A1" and value_input_option == "RAW"
        self.values = rows

    def resize(self, rows, cols):
        self.row_count, self.col_count = rows, cols


class FakeSpreadsheet:
    def __init__(self, tabs):
        self.ws = {t: FakeWS(t, i, v) for i, (t, v) in enumerate(tabs.items())}
        self.batches = []

    def worksheets(self):
        return list(self.ws.values())

    def worksheet(self, name):
        if name not in self.ws:
            raise type("WorksheetNotFound", (Exception,), {})(name)
        return self.ws[name]

    def add_worksheet(self, title, rows, cols):
        self.ws[title] = FakeWS(title, len(self.ws) + 100)
        return self.ws[title]

    def batch_update(self, body):
        self.batches.append(body)


class FakeClient:
    def __init__(self, sh):
        self.sh, self.opened = sh, None

    def open_by_url(self, url):
        self.opened = url
        return self.sh

    def open_by_key(self, key):
        self.opened = key
        return self.sh


def _as_google(path):
    """The xlsx fixture as Google would return it: numbers raw, dates as serial numbers."""
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            conv = []
            for v in row:
                if hasattr(v, "year"):
                    v = (pd.Timestamp(v) - pd.Timestamp("1899-12-30")) / pd.Timedelta(days=1)
                conv.append("" if v is None else v)
            rows.append(conv)
        out[ws.title] = rows
    return out


def test_google_backend_end_to_end_with_fake_api(cfg, sheet):
    sh = FakeSpreadsheet(_as_google(sheet))
    be = GoogleSheetsBackend("https://docs.google.com/spreadsheets/d/abc123/edit", "unused.json", client=FakeClient(sh))
    reps = bootstrap(cfg, catch_up=["Preserved Week 2 Line Export"], as_of=NOW, backend=be)
    assert [r.version for r in reps] == [1, 2, 3] and reps[1].counts["lines_refreshed"] == 32
    grid = sh.ws["Probability Grid (Latest)"].values
    assert grid[0][:3] == ["Team", "Week 1", "Week 2"] and len(grid) >= 33
    assert all(isinstance(v, (str, int, float)) for r in grid for v in r)            # JSON-safe for the API
    reqs = [r for b in sh.batches for r in b["requests"]]
    assert any("updateSheetProperties" in r for r in reqs) and any("numberFormat" in str(r) for r in reqs)
    for name in ("Game Lines", "Preserved Week 1 Full Season Li", "Preseason Win Totals"):
        assert sh.ws[name].cleared == 0                                               # inputs never touched


def test_google_format_requests_compress_fill_runs():
    spec = TabSpec(name="x", values=[["a"] * 5] * 2, fills={(1, 0): "#D9D9D9", (1, 1): "#D9D9D9", (1, 3): "#FCE4D6"})
    reqs = GoogleSheetsBackend.format_requests(7, spec, 2, 5)
    fills = [r["repeatCell"]["range"] for r in reqs if "repeatCell" in r and "backgroundColor" in str(r)]
    assert fills == [{"sheetId": 7, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 0, "endColumnIndex": 2},
                     {"sheetId": 7, "startRowIndex": 1, "endRowIndex": 2, "startColumnIndex": 3, "endColumnIndex": 4}]
